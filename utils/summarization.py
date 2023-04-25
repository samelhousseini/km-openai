import os
import pickle
import numpy as np
import pandas as pd
import urllib
from datetime import datetime, timedelta
import logging
import copy
import uuid
import json
import openpyxl
import time

from langchain import OpenAI, PromptTemplate, LLMChain
from langchain.text_splitter import CharacterTextSplitter
from langchain.chains.mapreduce import MapReduceChain
from langchain.prompts import PromptTemplate
from langchain.text_splitter import TokenTextSplitter, TextSplitter
from langchain.chains.summarize import load_summarize_chain
from langchain.docstore.document import Document
from langchain.callbacks.base import CallbackManager

from utils import openai_helpers
from utils import helpers
from utils import fr_helpers


MAX_OUTPUT_TOKENS   = int(os.environ["MAX_OUTPUT_TOKENS"])
CHOSEN_COMP_MODEL   = os.environ['CHOSEN_COMP_MODEL']



## Use
# """
# from utils import langchain_agent
# from utils import openai_helpers
# from utils import fr_helpers
# from utils import summarization
# folder = './docs_to_summarize'
# ref_summ_df = summarization.summarize_folder(folder, mode='refine', verbose=False)
# mp_summ_df  = summarization.summarize_folder(folder, mode='map_reduce', verbose=False)
# """





mapreduce_prompt_template = """The maximum output is about 500 to 750 tokens, so make sure to take advantage of this to the maximum.\n
Write an elaborate summary of 3 paragraphs of the following:


{text}


SUMMARY:"""


refine_prompt_template = """Write an elaborate summary of 3 paragraphs of the following:

{text}

"""

refine_template = (
    "Your job is to produce a final summary of 3 paragraphs that is elaborate and rich in details.\n" 
    "The maximum output is about 500 to 750 tokens, so make sure to take advantage of this to the maximum.\n"
    "We have provided an existing summary up to a certain point: {existing_answer}\n"
    "We have the opportunity to refine the existing summary."
    "(only if needed) with some more context below.\n"
    "------------\n"
    "{text}\n"
    "------------\n"
    "Given the new context, refine the original summary."
    "If the context isn't useful, return the original summary."
)



def chunk_doc(all_text, mode='refine', model=CHOSEN_COMP_MODEL, max_output_tokens=MAX_OUTPUT_TOKENS, chunk_overlap=500):

    enc_name = openai_helpers.get_encoding_name(model)
    enc = openai_helpers.get_encoder(model)

    max_tokens = openai_helpers.get_model_max_tokens(model)

    if mode == 'refine':
        max_tokens = max_tokens - len(enc.encode(refine_prompt_template)) - len(enc.encode(refine_template)) - 2*MAX_OUTPUT_TOKENS - chunk_overlap
    elif mode == 'map_reduce':
        max_tokens = max_tokens - len(enc.encode(mapreduce_prompt_template)) - MAX_OUTPUT_TOKENS - chunk_overlap
    else:
        raise Exception('Invalid mode')

    text_splitter = TokenTextSplitter(encoding_name=enc_name, chunk_size = max_tokens, chunk_overlap=chunk_overlap)
    
    texts = text_splitter.split_text(all_text)
    docs = [Document(page_content=t) for t in texts]

    enc = openai_helpers.get_encoder(CHOSEN_COMP_MODEL)

    l_arr = []
    for d in texts:
        l_arr.append(str(len(enc.encode(d))))

    print("Chunks Generated", len(docs), ' | max_tokens', max_tokens, " | Chunk Lengths:", ', '.join(l_arr))

    return docs


def clean_up_text(text):
    text = text.replace('....', '')
    return text



def get_refined_summarization(docs, model=CHOSEN_COMP_MODEL, max_output_tokens=MAX_OUTPUT_TOKENS, stream=False, callbacks=[]):

    PROMPT = PromptTemplate(template=refine_prompt_template, input_variables=["text"])
    refine_prompt = PromptTemplate(input_variables=["existing_answer", "text"],template=refine_template)

    llm = helpers.get_llm(model, temperature=0, max_output_tokens=max_output_tokens, stream=stream, callbacks=callbacks)

    chain = load_summarize_chain(llm, chain_type="refine",  question_prompt=PROMPT, refine_prompt=refine_prompt, return_intermediate_steps=True)
    summ = chain({"input_documents": docs}, return_only_outputs=True)
    
    return summ


def get_mapreduced_summarization(docs, model=CHOSEN_COMP_MODEL, max_output_tokens=MAX_OUTPUT_TOKENS, stream=False, callbacks=[]):

    PROMPT = PromptTemplate(template=mapreduce_prompt_template, input_variables=["text"])

    llm = helpers.get_llm(model, temperature=0, max_output_tokens=max_output_tokens, stream=stream, callbacks=callbacks)

    chain = load_summarize_chain(llm, chain_type="map_reduce", map_prompt=PROMPT, combine_prompt=PROMPT, return_intermediate_steps=True)
    summ = chain({"input_documents": docs}, return_only_outputs=True)
    
    return summ




def read_document(path, verbose = False):
    if verbose: print(f"Reading {path}")
    
    all_text = ''
    ext = os.path.splitext(path)[1]

    if ext == '.xlsx':
        dataframe = openpyxl.load_workbook(path, data_only=True)
        sheets = [s for s in dataframe.sheetnames if 'HiddenCache' not in s]
        for sheet in sheets:
            print('sheet', sheet)
            all_text += pd.read_excel(path, sheet_name=sheets[0]).to_string(na_rep='') + '\n\n\n\n'
    elif ext == '.csv':
        return None
    elif ext == '.pdf':
        contents, kv_contents, dfs, t_contents = fr_helpers.fr_analyze_local_doc_with_dfs(path, verbose = verbose)
        all_text = ' '.join([kv_contents , contents ,  t_contents])
    else:
        return None
    
    all_text = clean_up_text(all_text)

    return all_text


def summarize_document(path, mode='refine', verbose = False):

    print(f"##########################\nStarting Processing {path} ...")
    start = time.time()
    text = read_document(path, verbose=verbose)
    if text is None: return None

    summ = summarize_text(text, mode=mode, verbose=verbose)
    end = time.time()

    summary = {
        'file': os.path.basename(path),
        'intermediate_steps': summ['intermediate_steps'],
        'summary': summ['output_text'],
        'proc_time': end-start
    }

    print(f"Done Processing {path} in {end-start} seconds\n##########################\n")
    return summary 


def summarize_text(text, mode='refine', verbose = False):    
    docs = chunk_doc(text, mode=mode)

    if mode == 'refine':
        summ = get_refined_summarization(docs)
    elif mode == 'map_reduce':
        summ = get_mapreduced_summarization(docs)
    else:
        raise Exception("Invalid mode")

    return summ



def summarize_folder(folder, mode='refine', save_to_csv=True, save_to_pkl=True, verbose = False):
    files = os.listdir(folder)
    print(f"Files in folder {len(files)}")
    pkl_file = os.path.join(folder, f'summaries_{mode}.pkl')
    csv_file = os.path.join(folder, f'summaries_{mode}.csv')

    if os.path.exists(csv_file):
        summ_df = pd.read_csv(csv_file)
    else:
        summ_df = pd.DataFrame(columns=['file', 'intermediate_steps', 'summary', 'proc_time'])

    processed_files = list(summ_df['file'])
    print(f"List of already processed files {processed_files}")
     
    for f in files:        
        path = os.path.join(folder, f)
        if f in processed_files: continue
        
        summary = summarize_document(path, mode=mode, verbose=verbose)
        if summary is None: continue
        summ_df = pd.concat([summ_df, pd.DataFrame([summary])], ignore_index=True)

        if save_to_csv: summ_df.to_csv(csv_file)
        if save_to_pkl: summ_df.to_pickle(pkl_file)

    return summ_df